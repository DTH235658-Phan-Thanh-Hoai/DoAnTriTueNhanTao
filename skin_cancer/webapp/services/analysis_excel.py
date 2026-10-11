"""Create a styled batch analysis workbook with embedded local images."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins
from PIL import Image as PillowImage
from PIL import UnidentifiedImageError

from .analysis_export_data import binary_label, latest, load_batch_graph
from .batch_upload import spreadsheet_safe

SHEET_NAME = "ẢNH TRONG BATCH"
HEADERS = (
    "ID",
    "ẢNH GỐC",
    "GRAD-CAM",
    "TÊN FILE",
    "TRẠNG THÁI",
    "DỰ ĐOÁN",
    "TIN CẬY",
    "PHÂN LOẠI",
    "XỬ LÝ (ms)",
)
GREEN = "2BA149"
PALE_GREEN = "EAF7EF"
WHITE = "FFFFFF"
GRAY = "667085"
THUMBNAIL_PX = 96
IMAGE_COLUMN_PX = 168
IMAGE_ROW_PX = 117


class AnalysisExcelService:
    def __init__(self, app):
        self.upload_root = Path(app.config["UPLOAD_FOLDER"]).resolve()
        self.gradcam_root = Path(
            app.config.get(
                "GRADCAM_FOLDER",
                Path(app.config["PROJECT_ROOT"]) / "instance" / "gradcam",
            )
        ).resolve()

    def render(self, batch):
        batch = load_batch_graph(batch.id)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        self._configure_sheet(sheet, batch.id)

        for index, image in enumerate(sorted(batch.images, key=lambda item: item.id)):
            self._write_image_row(sheet, 6 + index, image, index)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _configure_sheet(sheet, batch_id):
        sheet.merge_cells("A1:I1")
        sheet.merge_cells("A2:I2")
        sheet.merge_cells("A3:I3")
        sheet["A1"] = "DERMASCAN AI — BATCH ANALYSIS"
        sheet["A2"] = f"BATCH #{batch_id} — IMAGE ANALYSIS REPORT"
        sheet["A3"] = "Kết quả AI chỉ mang tính hỗ trợ, không thay thế chẩn đoán y khoa."

        for row in (1, 2):
            cell = sheet.cell(row, 1)
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(
                name="Calibri",
                size=18 if row == 1 else 11,
                bold=True,
                color=WHITE,
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet["A3"].fill = PatternFill("solid", fgColor=PALE_GREEN)
        sheet["A3"].font = Font(name="Calibri", size=10, italic=True, color=GRAY)
        sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[2].height = 22
        sheet.row_dimensions[3].height = 22

        thin = Side(style="thin", color="D0D5DD")
        for column, label in enumerate(HEADERS, 1):
            cell = sheet.cell(5, column, label)
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        widths = (9, 24, 24, 30, 18, 17, 17, 18, 17)
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + column)].width = width
        sheet.row_dimensions[5].height = 34
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = "A5:I5"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:5"
        sheet.page_margins = PageMargins(
            left=0.3,
            right=0.3,
            top=0.5,
            bottom=0.5,
            header=0.2,
            footer=0.2,
        )

    def _write_image_row(self, sheet, row_number, image, index):
        analysis = latest(image.analyses)
        result = latest(analysis.results) if analysis is not None else None
        fill_color = WHITE if index % 2 == 0 else PALE_GREEN
        thin = Side(style="thin", color="D0D5DD")

        values = (
            image.id,
            "Không có ảnh gốc",
            self._gradcam_placeholder(analysis),
            image.original_filename,
            image.status,
            result.predicted_class if result is not None else None,
            result.predicted_probability if result is not None else None,
            binary_label(result.binary_result) if result is not None else "Chưa phân tích",
            result.inference_time_ms if result is not None else None,
        )
        for column, value in enumerate(values, 1):
            cell = sheet.cell(
                row_number,
                column,
                self._excel_safe(value),
            )
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(name="Calibri", size=10, color="1D2939")
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        sheet.cell(row_number, 7).number_format = "0.00%"
        sheet.cell(row_number, 9).number_format = "0.00"
        sheet.row_dimensions[row_number].height = 88
        self._style_status(sheet.cell(row_number, 5), image.status)

        original_path = self._allowed_file(image.image_path, self.upload_root)
        if original_path is not None:
            if self._embed_image(sheet, original_path, row_number, 2):
                sheet.cell(row_number, 2).value = None
            else:
                sheet.cell(row_number, 2).value = "Không thể đọc ảnh gốc"

        if analysis is not None and analysis.gradcam_path:
            gradcam_path = self._allowed_file(
                analysis.gradcam_path,
                self.gradcam_root,
            )
            if gradcam_path is not None:
                if self._embed_image(sheet, gradcam_path, row_number, 3):
                    sheet.cell(row_number, 3).value = None
                else:
                    sheet.cell(row_number, 3).value = "Không thể đọc Grad-CAM"

    @staticmethod
    def _excel_safe(value):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return spreadsheet_safe(value)

    @staticmethod
    def _gradcam_placeholder(analysis):
        if analysis is None or analysis.status in {"PENDING", "PROCESSING"}:
            return "Chưa phân tích"
        return "Không có Grad-CAM"

    @staticmethod
    def _style_status(cell, status):
        colors = {
            "COMPLETED": ("D1FADF", "067647"),
            "UPLOADED": ("FEF0C7", "B54708"),
            "PROCESSING": ("D1E9FF", "175CD3"),
            "FAILED": ("FEE4E2", "B42318"),
        }
        fill, font = colors.get(status, ("F2F4F7", "344054"))
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Calibri", size=10, bold=True, color=font)

    @staticmethod
    def _allowed_file(raw_path, root):
        if not raw_path:
            return None
        path = Path(raw_path).resolve()
        if root not in path.parents or not path.is_file():
            return None
        return path

    @staticmethod
    def _embed_image(sheet, path, row_number, column_number):
        try:
            payload = path.read_bytes()
            with PillowImage.open(BytesIO(payload)) as source:
                source.load()
                source.thumbnail(
                    (THUMBNAIL_PX, THUMBNAIL_PX),
                    PillowImage.Resampling.LANCZOS,
                )
                if source.mode not in {"1", "L", "LA", "P", "RGB", "RGBA"}:
                    target_mode = "RGBA" if "A" in source.getbands() else "RGB"
                    source = source.convert(target_mode)
                width, height = source.size
                normalized = BytesIO()
                source.save(normalized, format="PNG")
            if width <= 0 or height <= 0:
                return False
            normalized.seek(0)
            display_width = width
            display_height = height
            image = ExcelImage(normalized)
            image.width = display_width
            image.height = display_height
            x_offset = max(0, (IMAGE_COLUMN_PX - display_width) // 2)
            y_offset = max(0, (IMAGE_ROW_PX - display_height) // 2)
            image.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=column_number - 1,
                    colOff=pixels_to_EMU(x_offset),
                    row=row_number - 1,
                    rowOff=pixels_to_EMU(y_offset),
                ),
                ext=XDRPositiveSize2D(
                    pixels_to_EMU(display_width),
                    pixels_to_EMU(display_height),
                ),
            )
            sheet.add_image(image)
            return True
        except (
            OSError,
            SyntaxError,
            UnidentifiedImageError,
            ValueError,
            PillowImage.DecompressionBombError,
        ):
            return False
